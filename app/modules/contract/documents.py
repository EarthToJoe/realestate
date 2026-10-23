"""The contract as PDF (WeasyPrint) and DOCX (python-docx), with the label on every page.

PDF: the label is printed in the top and bottom page margins of every page by CSS @page rules, and the attached
PCDS follows the rider. DOCX: the label is in the section's header and footer, which Word repeats on every page; the
parties can edit the DOCX in Word. Party names are inserted here only; the caller logs the access.
"""

from __future__ import annotations

import hashlib
import io
import threading
from collections import OrderedDict

from app.brokerage import brokerage
from app.disclosures import pcds
from app.models import Deal
from app.modules.contract import service
from app.web import templates


def contract_html(deal: Deal, st: dict, mapping: dict[str, str], pcds_answers: dict | None) -> str:
    return templates.get_template("contract/print.html").render(
        deal=deal, st=st, body=st["draft"], label=service.LABEL, names=mapping, fill=service.fill,
        rows=service.side_by_side(st["draft"]), memo=st.get("memorandum"), firm=brokerage(),
        signing_map=service.SIGNING_MAP, standard_source=service.STANDARD_SOURCE,
        form=pcds.form(), questions=pcds.questions, fields=pcds.fields, labels=pcds.CHOICE_LABELS,
        saved=pcds_answers or {}, demo_note=pcds.DEMO_NOTE, attach_pcds=bool(pcds_answers))


# Each version's PDF is built once and kept in memory (never on disk: it carries the parties' names), keyed by a hash
# of its exact HTML, so a change of names, signatures or disclosures builds a new one. The contract screen warms it in
# the background, so the meeting never waits on WeasyPrint (over 30 s on the loaded server).
_PDF_CACHE: OrderedDict[str, bytes] = OrderedDict()
_BUILDING: dict[str, threading.Event] = {}
_LOCK = threading.Lock()
CACHE_SIZE = 12


def _key(html: str) -> str:
    return hashlib.sha256(html.encode()).hexdigest()


def contract_pdf(html: str) -> bytes:
    key = _key(html)
    with _LOCK:
        if key in _PDF_CACHE:
            _PDF_CACHE.move_to_end(key)
            return _PDF_CACHE[key]
        ev = _BUILDING.get(key)
        mine = ev is None
        if mine:
            ev = _BUILDING[key] = threading.Event()
    if not mine:  # another request or the warm-up is building this exact PDF: wait for it
        ev.wait(timeout=180)
        with _LOCK:
            if key in _PDF_CACHE:
                return _PDF_CACHE[key]
    try:
        from weasyprint import HTML

        data = HTML(string=html).write_pdf()
        with _LOCK:
            _PDF_CACHE[key] = data
            while len(_PDF_CACHE) > CACHE_SIZE:
                _PDF_CACHE.popitem(last=False)
        return data
    finally:
        if mine:
            with _LOCK:
                _BUILDING.pop(key, None)
            ev.set()


def cached(html: str) -> bool:
    with _LOCK:
        return _key(html) in _PDF_CACHE or _key(html) in _BUILDING


def warm(html: str) -> None:
    """Build the PDF in the background if it is not built or building yet."""
    if cached(html):
        return
    threading.Thread(target=lambda: _quiet(html), daemon=True, name="contract-pdf-warm").start()


def _quiet(html: str) -> None:
    try:
        contract_pdf(html)
    except Exception:  # noqa: BLE001 (a warm-up failure only means the request builds it)
        pass


def _field(run, instr: str) -> None:
    """Insert a Word field (e.g. PAGE) into a run."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for kind, text in (("begin", None), (None, instr), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def contract_docx(deal: Deal, st: dict, mapping: dict[str, str]) -> bytes:
    from docx import Document as Docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    body = st["draft"]
    doc = Docx()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    sec = doc.sections[0]
    for part in (sec.header, sec.footer):
        p = part.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(service.LABEL)
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0xAA, 0x00, 0x00)
    p = sec.footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Page ")
    r.font.size = Pt(8)
    _field(p.add_run(), "PAGE")
    doc.add_heading(service.fill(body["title"], mapping), level=1)
    doc.add_paragraph(f"Premises: {deal.property.street}, {deal.property.city}, NY {deal.property.zip}. "
                      f"Seller: {mapping['[SELLER_1]']}. Buyer: {mapping['[BUYER_1]']}. "
                      f"Draft version {st['version']} ({st['model']}, {st['mode']}).")
    for c in body["sections"]:
        doc.add_heading(f"{c['number']}. {service.fill(c['heading'], mapping)}", level=3)
        doc.add_paragraph(service.fill(c["text"], mapping))
    doc.add_page_break()
    doc.add_heading("Rider", level=1)
    for c in body["rider"]:
        doc.add_heading(f"{c['number']}. {service.fill(c['heading'], mapping)}", level=3)
        doc.add_paragraph(service.fill(c["text"], mapping))
    doc.add_heading("Signatures (simulated e-signature, L34)", level=2)
    signed = st.get("signed_at")
    for who in ("[SELLER_1]", "[BUYER_1]"):
        doc.add_paragraph(f"{mapping[who]}: " + (f"signed electronically {signed[:10]}" if signed else
                                                  "________________________   Date: __________"))
    doc.add_page_break()
    doc.add_heading("Side by side with the standard form's structure", level=1)
    doc.add_paragraph(service.STANDARD_SOURCE)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for i, h in enumerate(("Standard paragraph", "This draft", "Difference")):
        t.rows[0].cells[i].text = h
    for row in service.side_by_side(body):
        cells = t.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = f"{row['n']}. {row['heading']}", row["draft_ref"], row["note"]
    doc.add_heading("Signing map", level=2)
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for i, h in enumerate(("Document", "Who", "How", "Law; when")):
        t.rows[0].cells[i].text = h
    for row in service.SIGNING_MAP:
        cells = t.add_row().cells
        cells[0].text, cells[1].text, cells[2].text, cells[3].text = (row["doc"], row["who"], row["how"],
                                                                      f"{row['law']}; {row['when']}")
    doc.add_paragraph("Attached: the Property Condition Disclosure Statement completed by the seller (see the PDF), "
                      "and the other disclosures listed on the contract screen.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
